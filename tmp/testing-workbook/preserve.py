"""Apply authored cells while preserving the original workbook's native parts."""
import copy
import json
import re
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).parent
source = openpyxl.load_workbook(ROOT / 'original.xlsx')
authored = openpyxl.load_workbook(ROOT / 'authored-library.xlsx')
plan = json.loads((ROOT / 'plan.json').read_text())
removed = set(plan['removed'])
library = [list(row) for row in authored['Test Library'].values]
by_id = {row[0]: row for row in library[1:]}
ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
ET.register_namespace('', ns['s'])
ET.register_namespace('r', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships')

def tag(name):
    return '{' + ns['s'] + '}' + name

def write_cell(cell, value, cached=None):
    for child in list(cell):
        if child.tag in [tag('v'), tag('f'), tag('is')]:
            cell.remove(child)
    cell.attrib.pop('t', None)
    if value is None:
        return
    if isinstance(value, str) and value.startswith('='):
        ET.SubElement(cell, tag('f')).text = value[1:]
        if cached is not None:
            if isinstance(cached, str):
                cell.set('t', 'str')
            ET.SubElement(cell, tag('v')).text = str(cached)
    elif isinstance(value, str):
        cell.set('t', 'inlineStr')
        text = ET.SubElement(ET.SubElement(cell, tag('is')), tag('t'))
        text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        text.text = value
    elif isinstance(value, bool):
        cell.set('t', 'b')
        ET.SubElement(cell, tag('v')).text = str(int(value))
    else:
        ET.SubElement(cell, tag('v')).text = str(value)

def translate(value, old_row, new_row, col):
    if isinstance(value, str) and value.startswith('='):
        return Translator(value, origin=f'{col}{old_row}').translate_formula(f'{col}{new_row}')
    return value

# Retain row IDs, recorded results, assignees, evidence and notes for surviving tests.
old_library_rows = {row[0].value: row[0].row for row in source['Test Library'].iter_rows(min_row=2)}
old_execution_rows = {row[2].value: row[0].row for row in source['Test Execution'].iter_rows(min_row=2)}
execution = [list(next(source['Test Execution'].values))]
execution_origins = [1]
for row_number, test in enumerate(library[1:], 2):
    test_id = test[0]
    old_row = old_execution_rows.get(test_id)
    if old_row:
        entry = [translate(c.value, old_row, row_number, c.column_letter) for c in source['Test Execution'][old_row]]
    else:
        entry = [f'BETARUN-{test_id[3:]}', 'Beta Release', test_id, None, None, None, None,
                 'Staging', 'Not Run', None, None, None, None, None, 'Not Required', None, None]
        for column, library_column in [(3, 2), (4, 3), (5, 5), (6, 6), (15, 14)]:
            entry[column] = f'=IFERROR(VLOOKUP($C{row_number},\'Test Library\'!$A:$P,{library_column},FALSE),"")'
    execution.append(entry)
    execution_origins.append(old_row or 2)

coverage = [list(r) for r in source['Coverage Matrix'].values]
coverage[14][0] = 'Manager Dashboard'
checklist = [list(r) for r in source['Release Checklist'].values if r[0] != 'RC-024']
check_origins = [r[0].row for r in source['Release Checklist'].iter_rows() if r[0].value != 'RC-024']
checklist += [
    ['RC-051', 'Performance', 'OnePlus Nord 4 cold launch, idle resume and core page navigation measured against documented budgets', None, 'Yes', 'Not Started', None, 'Record build, network, OS, browser and five-run timings.'],
    ['RC-052', 'Dashboard', 'Manager actions and team filters verified at 320, 348, 360 and 390 CSS pixels', None, 'Yes', 'Not Started', None, 'Deploy the mobile layout change before running these checks.'],
]
check_origins += [2, 2]
devices = [list(r) for r in source['Devices & Environments'].values]
devices += [
    ['ENV-18', 'Staging', 'Mobile', 'OnePlus Nord 4 (record CSS viewport)', 'Installed Chrome PWA', 'Record installed Android / OxygenOS version', 'Critical', 'Cold launch, Android process eviction, idle resume, core page loads', None, None, 'Not Tested', 'Record Chrome version, battery settings and Wi-Fi/mobile data conditions.'],
    ['ENV-19', 'Staging', 'Mobile', '320 and 348 CSS pixels wide', 'Chrome responsive mode plus real-device check', 'Android', 'High', 'Manager dashboard action row and zone/status filter row', None, None, 'Not Tested', 'Also compare 360 and 390 CSS pixels and desktop.'],
]
test_data = [list(r) for r in source['Test Data'].values if r[0] != 'TD-014']
data_origins = [r[0].row for r in source['Test Data'].iter_rows() if r[0].value != 'TD-014']
updates = {
    'Test Library': (library, [1] + [old_library_rows.get(r[0], 2) for r in library[1:]]),
    'Test Execution': (execution, execution_origins),
    'Coverage Matrix': (coverage, list(range(1, len(coverage)+1))),
    'Release Checklist': (checklist, check_origins),
    'Devices & Environments': (devices, list(range(1, 19)) + [2, 2]),
    'Test Data': (test_data, data_origins),
}

# Formula cache values reflect the edited inputs, while formulas remain editable.
status_counts = Counter(r[8] for r in execution[1:])
gates = sum(by_id[r[2]][13] == 'Yes' and r[8] != 'Passed' for r in execution[1:])
check_counts = Counter(r[5] for r in checklist[1:])
required_pending = sum(r[4] == 'Yes' and r[5] != 'Passed' for r in checklist[1:])
dashboard_cache = {'A6':len(library)-1, 'C6':status_counts['Passed'], 'A8':status_counts['Failed'],
                   'C8':status_counts['Blocked'], 'A10':sum(status_counts[x] for x in ['Passed','Failed','Blocked','In Progress']),
                   'C10':0, 'B14':gates, 'B16':required_pending, 'A19':'NO-GO'}
for row in range(5, 11):
    dashboard_cache[f'G{row}'] = status_counts[source['Dashboard'][f'F{row}'].value]
for row in range(14, 19):
    dashboard_cache[f'G{row}'] = check_counts[source['Dashboard'][f'F{row}'].value]

def cached_value(name, row_number, column):
    if name == 'Test Execution' and column in [3,4,5,6,15] and row_number > 1:
        return by_id[execution[row_number-1][2]][{3:1,4:2,5:4,6:5,15:13}[column]]
    if name != 'Coverage Matrix' or row_number == 1 or column == 0:
        return None
    area = coverage[row_number-1][0]
    cases = [r for r in library[1:] if r[1] == area]
    runs = [r for r in execution[1:] if by_id[r[2]][1] == area]
    counts = {
        1:len(cases), 2:sum(r[5]=='Critical' for r in cases), 3:sum(r[5]=='High' for r in cases),
        4:sum(r[5]=='Medium' for r in cases), 5:sum(r[5]=='Low' for r in cases),
        6:sum(r[3]=='Security' for r in cases), 7:sum(r[3]=='Integration' for r in cases),
        8:sum(r[3]=='Accessibility' for r in cases),
        9:sum(r[3] in ['Performance','Recovery','Reliability','Concurrency'] for r in cases),
        10:sum(r[12]!='Manual' for r in cases), 11:sum(r[13]=='Yes' for r in cases),
        12:sum(r[8]=='Passed' for r in runs),13:sum(r[8]=='Failed' for r in runs),
        14:sum(r[8]=='Blocked' for r in runs),15:sum(r[8]=='Not Run' for r in runs),
        16:0,17:'Not Ready' if any(by_id[r[2]][13]=='Yes' and r[8]!='Passed' for r in runs) else 'Ready',
    }
    return counts[column]

with zipfile.ZipFile(ROOT/'original.xlsx') as original:
    parts = {name:original.read(name) for name in original.namelist()}
for name, (values, origins) in updates.items():
    sheet_index = source.sheetnames.index(name)+1
    path = f'xl/worksheets/sheet{sheet_index}.xml'
    tree = ET.fromstring(parts[path])
    sheet_data = tree.find('s:sheetData', ns)
    original_rows = {int(r.attrib['r']):r for r in sheet_data}
    for row in list(sheet_data):
        sheet_data.remove(row)
    for row_number, (values_row, origin) in enumerate(zip(values, origins), 1):
        row = copy.deepcopy(original_rows[origin])
        row.set('r',str(row_number))
        old_cells = {re.sub(r'\d','',c.attrib['r']):c for c in row}
        for cell in list(row):
            row.remove(cell)
        for column, value in enumerate(values_row):
            letter = get_column_letter(column+1)
            cell = copy.deepcopy(old_cells.get(letter, ET.Element(tag('c'))))
            cell.set('r', f'{letter}{row_number}')
            write_cell(cell,value,cached_value(name,row_number,column))
            row.append(cell)
        if name == 'Test Library' and values_row[0] in plan['added']:
            row.set('ht','180')
            row.set('customHeight','1')
        sheet_data.append(row)
    old_end = source[name].max_row
    new_end = len(values)
    for el in tree.iter():
        for attr in ['ref','sqref']:
            if attr in el.attrib:
                el.set(attr,re.sub(r'(?<=[A-Z])'+str(old_end)+r'\b', str(new_end), el.attrib[attr]))
    parts[path] = ET.tostring(tree,encoding='utf-8',xml_declaration=True)
    for table in source[name].tables.values():
        table_path = next(p for p in parts if p.startswith('xl/tables/') and ET.fromstring(parts[p]).get('name') == table.name)
        table_xml = ET.fromstring(parts[table_path])
        new_ref = f'A1:{get_column_letter(len(values[0]))}{new_end}'
        table_xml.set('ref',new_ref)
        auto_filter = table_xml.find('s:autoFilter',ns)
        if auto_filter is not None:
            auto_filter.set('ref',new_ref)
        parts[table_path] = ET.tostring(table_xml,encoding='utf-8',xml_declaration=True)

readme = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
for cell in readme.findall('.//s:c',ns):
    if cell.attrib['r'] == 'A5':
        write_cell(cell, f'This workbook contains {plan["count"]} RocketRota-specific test cases covering the beta product, NFC clock-in security, billing, PWA behaviour, permissions, operations and production readiness. Updated 25 October 2026. Use it as the master QA library and record each release in Test Execution.')
parts['xl/worksheets/sheet1.xml'] = ET.tostring(readme,encoding='utf-8',xml_declaration=True)
dashboard = ET.fromstring(parts['xl/worksheets/sheet2.xml'])
for cell in dashboard.findall('.//s:c',ns):
    address = cell.attrib['r']
    if address in dashboard_cache:
        write_cell(cell,source['Dashboard'][address].value,dashboard_cache[address])
parts['xl/worksheets/sheet2.xml'] = ET.tostring(dashboard,encoding='utf-8',xml_declaration=True)
# Keep chart formulas and styling; refresh the status chart's numeric snapshot.
chart_ns = {'c':'http://schemas.openxmlformats.org/drawingml/2006/chart'}
for path in [p for p in parts if p.startswith('xl/charts/chart') and p.endswith('.xml')]:
    chart = ET.fromstring(parts[path])
    changed = False
    for ref in chart.findall('.//c:numRef',chart_ns):
        formula = ref.find('c:f',chart_ns)
        if formula is not None and '$G$5:$G$10' in (formula.text or ''):
            for point in ref.findall('.//c:pt',chart_ns):
                index = int(point.get('idx'))
                point.find('c:v',chart_ns).text = str(dashboard_cache[f'G{5+index}'])
            changed=True
    if changed:
        parts[path]=ET.tostring(chart,encoding='utf-8',xml_declaration=True)
workbook_xml = ET.fromstring(parts['xl/workbook.xml'])
calc = workbook_xml.find('s:calcPr',ns)
if calc is not None:
    calc.set('fullCalcOnLoad','1')
    calc.set('forceFullCalc','1')
parts['xl/workbook.xml'] = ET.tostring(workbook_xml,encoding='utf-8',xml_declaration=True)

output = ROOT/'RocketRota_Full_Testing_Workbook.xlsx'
with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as result:
    for path,data in parts.items():
        result.writestr(path,data)
print(json.dumps({'output':str(output),'removed':len(removed),'added':len(plan['added']),'total':plan['count']}))
